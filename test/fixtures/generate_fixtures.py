#!/usr/bin/env python3
"""Generate test fixture files for converter tests."""
import os

FIXTURES_DIR = os.path.dirname(os.path.abspath(__file__))


def generate_html():
    path = os.path.join(FIXTURES_DIR, "sample.html")
    with open(path, "w") as f:
        f.write("""<!DOCTYPE html>
<html>
<head><title>th2rag HTML test</title></head>
<body>
<h1>th2rag HTML fixture</h1>
<p>This is a test paragraph for HTML conversion in th2rag RAG pipeline.</p>
<p>The converter should extract this text successfully.</p>
</body>
</html>""")
    print(f"Generated: {path}")
    return path


def generate_markdown():
    path = os.path.join(FIXTURES_DIR, "sample.md")
    with open(path, "w") as f:
        f.write("""# th2rag Markdown fixture

This is a test paragraph for Markdown conversion in th2rag RAG pipeline.

## Section 2

The converter should extract this text successfully.

- Item 1
- Item 2
- Item 3
""")
    print(f"Generated: {path}")
    return path


def generate_image_with_text():
    """Generate a PNG image with rendered text using Pillow."""
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (400, 100), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((10, 30), "th2rag OCR test", fill=(0, 0, 0))

    path = os.path.join(FIXTURES_DIR, "sample_text.png")
    img.save(path)
    print(f"Generated: {path}")
    return path


def generate_simple_docx():
    """Generate a simple DOCX with text content."""
    try:
        from docx import Document
        doc = Document()
        doc.add_heading("th2rag DOCX fixture", 0)
        doc.add_paragraph("This is a test paragraph for DOCX conversion in th2rag RAG pipeline.")
        doc.add_paragraph("The converter should extract this text successfully.")
        path = os.path.join(FIXTURES_DIR, "sample_text.docx")
        doc.save(path)
        print(f"Generated: {path}")
        return path
    except ImportError:
        print("python-docx not available, skipping DOCX fixture")
        return None


if __name__ == "__main__":
    generate_html()
    generate_markdown()
    generate_image_with_text()
    generate_simple_docx()
    print("All fixtures generated.")
